"""Build full or partial annual workbooks from the embedded template."""

from datetime import date
from openpyxl import Workbook
from openpyxl.utils import get_column_letter
from openpyxl.workbook.workbook import Workbook as WorkbookType
from openpyxl.worksheet.worksheet import Worksheet
from .calendar_utils import get_month_weeks
from .excel_constants import (
    MONTHLY_CLOSE_END_COLUMN,
    MONTHLY_CLOSE_END_ROW,
    MONTHLY_CLOSE_START_COLUMN,
    MONTHLY_CLOSE_START_ROW,
    MONTH_NAMES,
    DAILY_EXCHANGE_RATES_SHEET,
    PROCESSED_DOCUMENTS_SHEET,
    WEEK_END_ROW,
    WEEK_START_ROW,
)
from .excel_copy import copy_block


def _get_day_start_column(week_start_column: int, *, include_labels: bool) -> int:
    """Get day start column."""
    return week_start_column + int(include_labels)


def _get_week_trim_counts(week: list[date]) -> tuple[int, int]:
    """Get week trim counts."""
    weekdays = [day.weekday() for day in week if day.weekday() < 6]
    if not weekdays:
        return (6, 0)
    leading_days = min(weekdays)
    trailing_days = 5 - max(weekdays)
    return (leading_days, trailing_days)


def _copy_week_days(
    template_sheet: Worksheet,
    target_sheet: Worksheet,
    week: list[date],
    target_start_column: int,
) -> int:
    """Copy week days."""
    leading_days, trailing_days = _get_week_trim_counts(week)
    day_count = 6 - leading_days - trailing_days
    if day_count == 0:
        return 0
    source_start_column = 2 + leading_days * 2
    source_end_column = 13 - trailing_days * 2
    for min_row, max_row, clear_values in (
        (WEEK_START_ROW, 3, False),
        (4, WEEK_END_ROW, True),
    ):
        copy_block(
            source_sheet=template_sheet,
            target_sheet=target_sheet,
            min_row=min_row,
            max_row=max_row,
            min_column=source_start_column,
            max_column=source_end_column,
            target_start_row=min_row,
            target_start_column=target_start_column,
            clear_values=clear_values,
        )
    return day_count


def _copy_week_closing(
    template_sheet: Worksheet, target_sheet: Worksheet, target_start_column: int
) -> None:
    """Copy week closing."""
    for min_row, max_row, clear_values in (
        (WEEK_START_ROW, 3, False),
        (4, WEEK_END_ROW, True),
    ):
        copy_block(
            source_sheet=template_sheet,
            target_sheet=target_sheet,
            min_row=min_row,
            max_row=max_row,
            min_column=14,
            max_column=15,
            target_start_row=min_row,
            target_start_column=target_start_column,
            clear_values=clear_values,
            clear_formulas=clear_values,
        )


def _set_week_dates(
    sheet: Worksheet, week: list[date], week_start_column: int, *, include_labels: bool
) -> None:
    """Set week dates."""
    days = [day for day in week if day.weekday() < 6]
    if not days:
        return
    day_start_column = _get_day_start_column(
        week_start_column=week_start_column, include_labels=include_labels
    )
    for day_index, day in enumerate(days):
        column = day_start_column + day_index * 2
        sheet.cell(row=2, column=column).value = day
    closing_column = day_start_column + len(days) * 2
    sheet.cell(row=2, column=closing_column).value = days[0].isocalendar().week


def _set_week_closing_formulas(
    template_sheet: Worksheet,
    target_sheet: Worksheet,
    day_start_column: int,
    day_count: int,
) -> None:
    """Set week closing formulas."""
    closing_column = day_start_column + day_count * 2
    for row in range(4, WEEK_END_ROW + 1):
        for currency_offset in (0, 1):
            source = template_sheet.cell(row=row, column=14 + currency_offset)
            if source.data_type != "f":
                continue
            references = [
                f"{get_column_letter(day_start_column + day_index * 2 + currency_offset)}{row}"
                for day_index in range(day_count)
            ]
            target_sheet.cell(
                row=row, column=closing_column + currency_offset
            ).value = "=" + "+".join(references)


def _create_week_block(
    template_sheet: Worksheet,
    target_sheet: Worksheet,
    week: list[date],
    week_start_column: int,
    *,
    include_labels: bool,
) -> int:
    """Create week block."""
    leading_days, trailing_days = _get_week_trim_counts(week)
    day_count = 6 - leading_days - trailing_days
    if day_count == 0:
        return 0
    if include_labels:
        copy_block(
            source_sheet=template_sheet,
            target_sheet=target_sheet,
            min_row=WEEK_START_ROW,
            max_row=WEEK_END_ROW,
            min_column=1,
            max_column=1,
            target_start_row=WEEK_START_ROW,
            target_start_column=week_start_column,
        )
    day_start_column = _get_day_start_column(
        week_start_column=week_start_column, include_labels=include_labels
    )
    _copy_week_days(
        template_sheet=template_sheet,
        target_sheet=target_sheet,
        week=week,
        target_start_column=day_start_column,
    )
    closing_column = day_start_column + day_count * 2
    _copy_week_closing(
        template_sheet=template_sheet,
        target_sheet=target_sheet,
        target_start_column=closing_column,
    )
    _set_week_closing_formulas(
        template_sheet=template_sheet,
        target_sheet=target_sheet,
        day_start_column=day_start_column,
        day_count=day_count,
    )
    _set_week_dates(
        sheet=target_sheet,
        week=week,
        week_start_column=week_start_column,
        include_labels=include_labels,
    )
    return int(include_labels) + day_count * 2 + 2


def _create_month_weeks(
    template_sheet: Worksheet, target_sheet: Worksheet, year: int, month: int
) -> list[int]:
    """Create month weeks."""
    next_column = 1
    include_labels = True
    closing_columns = []
    for week in get_month_weeks(year, month):
        block_width = _create_week_block(
            template_sheet=template_sheet,
            target_sheet=target_sheet,
            week=week,
            week_start_column=next_column,
            include_labels=include_labels,
        )
        if block_width == 0:
            continue
        closing_columns.append(next_column + block_width - 2)
        next_column += block_width
        include_labels = False
    return closing_columns


def _build_monthly_sum(
    closing_columns: list[int], source_row: int, *, currency_offset: int = 0
) -> str:
    """Build monthly sum."""
    references = [
        f"{get_column_letter(column + currency_offset)}{source_row}"
        for column in closing_columns
    ]
    return "(" + "+".join(references) + ")"


def _copy_monthly_close(template_sheet: Worksheet, target_sheet: Worksheet) -> None:
    """Copy monthly close."""
    copy_block(
        source_sheet=template_sheet,
        target_sheet=target_sheet,
        min_row=MONTHLY_CLOSE_START_ROW,
        max_row=MONTHLY_CLOSE_END_ROW,
        min_column=MONTHLY_CLOSE_START_COLUMN,
        max_column=MONTHLY_CLOSE_START_COLUMN,
        target_start_row=MONTHLY_CLOSE_START_ROW,
        target_start_column=MONTHLY_CLOSE_START_COLUMN,
    )
    for min_row, max_row, clear_values in (
        (MONTHLY_CLOSE_START_ROW, 45, False),
        (46, MONTHLY_CLOSE_END_ROW, True),
    ):
        copy_block(
            source_sheet=template_sheet,
            target_sheet=target_sheet,
            min_row=min_row,
            max_row=max_row,
            min_column=MONTHLY_CLOSE_START_COLUMN + 1,
            max_column=MONTHLY_CLOSE_END_COLUMN,
            target_start_row=min_row,
            target_start_column=MONTHLY_CLOSE_START_COLUMN + 1,
            clear_values=clear_values,
        )


def _set_monthly_basic_formulas(sheet: Worksheet, closing_columns: list[int]) -> None:
    """Set monthly basic formulas."""
    row_mapping = {
        46: 12,
        47: 13,
        48: 14,
        50: 16,
        51: 17,
        52: 18,
        54: 20,
        55: 21,
        72: 40,
        73: 32,
        75: 22,
        76: 23,
        77: 24,
    }
    for target_row, source_row in row_mapping.items():
        sheet.cell(row=target_row, column=2).value = "=" + _build_monthly_sum(
            closing_columns, source_row
        )


def _set_monthly_currency_formulas(
    sheet: Worksheet, closing_columns: list[int], exchange_rate_reference: str
) -> None:
    """Set monthly currency formulas."""
    row_mapping = {
        58: 26,
        60: 28,
        61: 29,
        62: 30,
        63: 31,
        65: 34,
        66: 35,
        67: 36,
        68: 37,
        70: 38,
        71: 39,
    }
    for target_row, source_row in row_mapping.items():
        huf_sum = _build_monthly_sum(closing_columns, source_row)
        eur_sum = _build_monthly_sum(closing_columns, source_row, currency_offset=1)
        sheet.cell(row=target_row, column=2).value = (
            f"={huf_sum}+({eur_sum}*{exchange_rate_reference})"
        )


def _set_monthly_purchase_formula(sheet: Worksheet, closing_columns: list[int]) -> None:
    """Set monthly purchase formula."""
    purchases = _build_monthly_sum(closing_columns, 27)
    sheet["B59"] = f"={purchases}-(B61+B62+B63+B66+B67+B68)"


def _create_monthly_close(
    template_sheet: Worksheet,
    target_sheet: Worksheet,
    closing_columns: list[int],
    month: int,
    exchange_rate_reference: str,
) -> None:
    """Create monthly close."""
    _copy_monthly_close(template_sheet=template_sheet, target_sheet=target_sheet)
    target_sheet["B43"] = "Monthly total"
    target_sheet["B44"] = MONTH_NAMES[month - 1]
    target_sheet["B45"] = "HUF"
    _set_monthly_basic_formulas(sheet=target_sheet, closing_columns=closing_columns)
    _set_monthly_currency_formulas(
        sheet=target_sheet,
        closing_columns=closing_columns,
        exchange_rate_reference=exchange_rate_reference,
    )
    _set_monthly_purchase_formula(sheet=target_sheet, closing_columns=closing_columns)


def _create_month_sheet(
    workbook: WorkbookType,
    template_sheet: Worksheet,
    year: int,
    month: int,
    exchange_rate_reference: str,
) -> Worksheet:
    """Create month sheet."""
    month_name = MONTH_NAMES[month - 1]
    sheet = workbook.create_sheet(title=month_name)
    closing_columns = _create_month_weeks(
        template_sheet=template_sheet, target_sheet=sheet, year=year, month=month
    )
    _create_monthly_close(
        template_sheet=template_sheet,
        target_sheet=sheet,
        closing_columns=closing_columns,
        month=month,
        exchange_rate_reference=exchange_rate_reference,
    )

    sheet.freeze_panes = sheet["D4"]

    return sheet


def _create_summary_sheet(workbook: WorkbookType, template_sheet: Worksheet) -> None:
    """Create summary sheet."""
    sheet = workbook.create_sheet("Summary", index=0)
    copy_block(
        source_sheet=template_sheet,
        target_sheet=sheet,
        min_row=MONTHLY_CLOSE_START_ROW,
        max_row=MONTHLY_CLOSE_END_ROW,
        min_column=1,
        max_column=1,
        target_start_row=1,
        target_start_column=1,
    )


def _create_summary_month(workbook: WorkbookType, month: int) -> None:
    """Create summary month."""
    month_name = MONTH_NAMES[month - 1]
    month_sheet = workbook[month_name]
    summary = workbook["Summary"]
    target_column = month + 1
    copy_block(
        source_sheet=month_sheet,
        target_sheet=summary,
        min_row=MONTHLY_CLOSE_START_ROW,
        max_row=MONTHLY_CLOSE_END_ROW,
        min_column=2,
        max_column=2,
        target_start_row=1,
        target_start_column=target_column,
        clear_formulas=True,
    )
    summary.column_dimensions[get_column_letter(target_column)].width = 22
    for source_row in range(46, MONTHLY_CLOSE_END_ROW + 1):
        target_row = source_row - MONTHLY_CLOSE_START_ROW + 1
        label = month_sheet.cell(row=source_row, column=1).value
        value = month_sheet.cell(row=source_row, column=2).value
        if label is None and value is None:
            continue
        reference = f"'{month_name}'!B{source_row}"
        summary.cell(row=target_row, column=target_column).value = (
            f'=IF({reference}="","",{reference})'
        )


def _create_exchange_rate_sheet(workbook: WorkbookType) -> None:
    """Create exchange rate sheet."""
    sheet = workbook.create_sheet("Exchange_Rates")
    sheet.append(["Month", "EUR/HUF rate", "Rate source date"])
    for month_name in MONTH_NAMES:
        sheet.append([month_name, 0, None])
    sheet.column_dimensions["A"].width = 18
    sheet.column_dimensions["B"].width = 22
    sheet.column_dimensions["C"].width = 18
    for row in range(2, 14):
        sheet.cell(row=row, column=2).number_format = "0.00"
        sheet.cell(row=row, column=3).number_format = "yyyy-mm-dd"


def _create_daily_exchange_rates_sheet(workbook: WorkbookType) -> None:
    """Create daily exchange rates sheet."""
    sheet = workbook.create_sheet(DAILY_EXCHANGE_RATES_SHEET)
    sheet.append(["Application date", "Rate source date", "EUR/HUF rate"])
    for column in ("A", "B", "C"):
        sheet.column_dimensions[column].width = 22
    sheet.sheet_state = "veryHidden"


def _create_processed_documents_sheet(workbook: WorkbookType) -> None:
    """Create processed documents sheet."""
    sheet = workbook.create_sheet(PROCESSED_DOCUMENTS_SHEET)
    sheet.append(["Year", "Document number"])
    sheet.column_dimensions["A"].width = 12
    sheet.column_dimensions["B"].width = 30
    sheet.sheet_state = "veryHidden"


def create_year_workbook(template_sheet: Worksheet, year: int) -> WorkbookType:
    """Create all twelve monthly sheets, summary, rates and hidden registers."""
    workbook = Workbook()
    workbook.remove(workbook.active)
    _create_exchange_rate_sheet(workbook)
    _create_summary_sheet(workbook, template_sheet)
    for month in range(1, 13):
        _create_month_sheet(
            workbook=workbook,
            template_sheet=template_sheet,
            year=year,
            month=month,
            exchange_rate_reference=f"'Exchange_Rates'!$B${month + 1}",
        )
        _create_summary_month(workbook, month)
    _create_daily_exchange_rates_sheet(workbook)
    _create_processed_documents_sheet(workbook)
    return workbook

def create_partial_year_workbook(template_sheet: Worksheet, year: int, months: list[int]) -> WorkbookType:
    """Create a reference workbook for selected months, keeping summary positions."""
    workbook = Workbook()
    workbook.remove(workbook.active)
    _create_exchange_rate_sheet(workbook)
    _create_summary_sheet(workbook, template_sheet)
    for month in months:
        _create_month_sheet(
            workbook=workbook,
            template_sheet=template_sheet,
            year=year,
            month=month,
            exchange_rate_reference=f"'Exchange_Rates'!$B${month + 1}",
        )
        _create_summary_month(workbook, month)
    _create_daily_exchange_rates_sheet(workbook)
    _create_processed_documents_sheet(workbook)
    return workbook