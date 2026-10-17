"""Store daily exchange rates and update the monthly exchange rate sheet."""

from datetime import date, datetime

from openpyxl.workbook.workbook import Workbook as WorkbookType

from .excel_constants import DAILY_EXCHANGE_RATES_SHEET
from .models import DailyExchangeRate

def _write_daily_exchange_rates(
    workbook: WorkbookType,
    year: int,
    exchange_rates: dict[date, DailyExchangeRate],
) -> None:
    """Write daily exchange rates."""
    sheet = workbook[DAILY_EXCHANGE_RATES_SHEET]

    stored_dates = {
        value.date() if isinstance(value, datetime) else value
        for (value,) in sheet.iter_rows(
            min_row=2,
            max_col=1,
            values_only=True,
        )
        if value is not None
    }

    for day, record in sorted(exchange_rates.items()):
        if day.year != year or day in stored_dates:
            continue

        row = sheet.max_row + 1

        application_cell = sheet.cell(row=row, column=1)
        application_cell.value = record.application_date
        application_cell.number_format = "yyyy-mm-dd"

        source_cell = sheet.cell(row=row, column=2)
        source_cell.value = record.source_date
        source_cell.number_format = "yyyy-mm-dd"

        rate_cell = sheet.cell(row=row, column=3)
        rate_cell.value = float(record.rate)
        rate_cell.number_format = "0.00"

def _write_monthly_exchange_rate(
    workbook: WorkbookType,
    year: int,
    exchange_rates: dict[date, DailyExchangeRate],
) -> None:
    """Write monthly exchange rate."""
    today = date.today()

    if year != today.year:
        return

    record = exchange_rates[today]
    sheet = workbook["Exchange_Rates"]
    row = today.month + 1

    sheet.cell(row=row, column=2).value = float(record.rate)

    source_cell = sheet.cell(row=row, column=3)
    source_cell.value = record.source_date
    source_cell.number_format = "yyyy-mm-dd"

def write_exchange_rates(
    workbook: WorkbookType,
    year: int,
    exchange_rates: dict[date, DailyExchangeRate],
) -> WorkbookType:
    """Store daily rates and update the current monthly rate in the workbook."""
    _write_daily_exchange_rates(
        workbook=workbook,
        year=year,
        exchange_rates=exchange_rates,
    )

    _write_monthly_exchange_rate(
        workbook=workbook,
        year=year,
        exchange_rates=exchange_rates,
    )

    return workbook