"""Load, validate and atomically save annual Excel reports."""

from pathlib import Path
from tempfile import NamedTemporaryFile
from time import perf_counter
from zipfile import BadZipFile
from xml.etree.ElementTree import ParseError
from datetime import date, datetime
from math import isfinite, isclose
from copy import copy
import warnings

from openpyxl import load_workbook
from openpyxl.workbook.workbook import Workbook as WorkbookType
from openpyxl.worksheet.worksheet import Worksheet
from openpyxl.utils import get_column_letter, column_index_from_string
from openpyxl.utils.exceptions import InvalidFileException
from openpyxl.cell.cell import MergedCell
from openpyxl.formula import Tokenizer
from openpyxl.styles import Color

from .excel_create_workbook import create_partial_year_workbook, create_year_workbook
from .excel_constants import (
    MONTH_NAMES,
    DAILY_EXCHANGE_RATES_SHEET,
    PROCESSED_DOCUMENTS_SHEET,
)
from .calendar_utils import get_month_weeks
from .template_creator import create_template_workbook

class WorkbookValidationError(ValueError):
    """Workbook validation error."""

class WorkbookValidationWarning(Warning):
    """Workbook validation warning."""

class WorkbookFileAccessError(OSError):
    """Workbook file access error."""

class WorkbookSaveError(OSError):
    """Workbook save error."""

def _get_validation_columns(
    sheet: Worksheet,
    months_name: list[str],
) -> list[int]:
    """Get validation columns."""
    if sheet.title == "Summary":
        return [
            1,
            *sorted(
                MONTH_NAMES.index(month_name) + 2
                for month_name in months_name
            ),
        ]

    return list(range(1, sheet.max_column + 1))

def _validate_required_sheets(
    workbook: WorkbookType,
    path: Path,
    months_name: list[str]
) -> None:
    """Validate required sheets."""

    required_sheets = [
        *months_name,
        "Summary",
        "Exchange_Rates",
        DAILY_EXCHANGE_RATES_SHEET,
        PROCESSED_DOCUMENTS_SHEET,
    ]

    missing_sheets = [
        name
        for name in required_sheets
        if name not in workbook.sheetnames
    ]

    if missing_sheets:
        missing_names = "\n".join(
            f"  - {name}" for name in missing_sheets
        )

        raise WorkbookValidationError(
            f"Missing or renamed worksheets.\n"
            f"File: {path}\n"
            f"Expected worksheet names not found:\n"
            f"{missing_names}\n"
            "Check worksheet names. Rename modified sheets back "
            "and restore deleted sheets."
        )

def _validate_header_cell(
    sheet: Worksheet,
    row: int,
    column: int,
    expected_value: str | int | date | None,
    path: Path,
) -> None:
    """Validate header cell."""
    cell = sheet.cell(row=row, column=column)
    actual_value = cell.value

    if isinstance(expected_value, date) and isinstance(actual_value, datetime):
        actual_value = actual_value.date()

    if actual_value != expected_value:
        raise WorkbookValidationError(
            f"Invalid header or displaced table section.\n"
            f"File: {path}\n"
            f"Worksheet: {sheet.title}\n"
            f"Cell: {cell.coordinate}\n"
            f"Current value: {cell.value!r}\n"
            f"Expected value: {expected_value!r}\n"
            "Restore the specified cell value. "
            "If rows or columns were inserted or deleted, "
            "restore the original table layout as well."
        )

def _validate_month_headers(
    sheet: Worksheet,
    year: int,
    month_name: str,
    path: Path,
) -> None:
    """Validate month headers."""
    month = MONTH_NAMES.index(month_name) + 1
    weekday_names = (
        "Monday",
        "Tuesday",
        "Wednesday",
        "Thursday",
        "Friday",
        "Saturday",
    )

    expected_cells = {
        (1, 1): "Date:",
        (3, 1): "Currency:",
        (43, 1): "Date:",
        (43, 2): "Monthly total",
        (44, 2): MONTH_NAMES[month - 1],
        (45, 1): "Currency:",
        (45, 2): "HUF",
    }

    next_column = 2

    for week in get_month_weeks(year, month):
        days = [day for day in week if day.weekday() < 6]

        if not days:
            continue

        for day in days:
            expected_cells[(1, next_column)] = weekday_names[day.weekday()]
            expected_cells[(2, next_column)] = day
            expected_cells[(3, next_column)] = "HUF"
            next_column += 2

        expected_cells[(1, next_column)] = "Weekly total"
        expected_cells[(2, next_column)] = days[0].isocalendar().week
        expected_cells[(3, next_column)] = "HUF"
        next_column += 2

    # Empty header cells must also match the reference.
    for row in (1, 2, 3):
        for column in range(1, next_column):
            expected_cells.setdefault((row, column), None)

    for row in (43, 44, 45):
        for column in range(1, 4):
            expected_cells.setdefault((row, column), None)

    for (row, column), expected_value in expected_cells.items():
        _validate_header_cell(
            sheet=sheet,
            row=row,
            column=column,
            expected_value=expected_value,
            path=path,
        )

def _validate_other_sheet_headers(
    workbook: WorkbookType,
    path: Path,
) -> None:
    """Validate other sheet headers."""
    summary = workbook["Summary"]

    for row, expected_value in (
        (1, "Date:"),
        (2, None),
        (3, "Currency:"),
    ):
        _validate_header_cell(
            sheet=summary,
            row=row,
            column=1,
            expected_value=expected_value,
            path=path,
        )

    for column, month_name in enumerate(MONTH_NAMES, start=2):
        for row, expected_value in (
            (1, "Monthly total"),
            (2, month_name),
            (3, "HUF")
        ):
            _validate_header_cell(
                sheet=summary,
                row=row,
                column=column,
                expected_value=expected_value,
                path=path,
            )

    headers = {
        "Exchange_Rates": (
            "Month",
            "EUR/HUF rate",
            "Rate source date",
        ),
        DAILY_EXCHANGE_RATES_SHEET: (
            "Application date",
            "Rate source date",
            "EUR/HUF rate",
        ),
        PROCESSED_DOCUMENTS_SHEET: (
            "Year",
            "Document number",
        ),
    }

    for sheet_name, column_headers in headers.items():
        for column, expected_value in enumerate(column_headers, start=1):
            _validate_header_cell(
                sheet=workbook[sheet_name],
                row=1,
                column=column,
                expected_value=expected_value,
                path=path,
            )

    euro = workbook["Exchange_Rates"]

    for row, month_name in enumerate(MONTH_NAMES, start=2):
        _validate_header_cell(
            sheet=euro,
            row=row,
            column=1,
            expected_value=month_name,
            path=path,
        )

def _validate_merged_ranges(
    workbook: WorkbookType,
    reference: WorkbookType,
    path: Path,   
) -> None:
    """Validate merged ranges."""
    for reference_sheet in reference:
        sheet = workbook[reference_sheet.title]

        expected_ranges = {
            str(merged_range)
            for merged_range in reference_sheet.merged_cells.ranges
        }

        actual_ranges = {
            str(merged_range)
            for merged_range in sheet.merged_cells.ranges
        }

        missing_ranges = sorted(expected_ranges - actual_ranges)
        unexpected_ranges = sorted(actual_ranges - expected_ranges)

        if not missing_ranges and not unexpected_ranges:
            continue

        details = []

        if missing_ranges:
            details.append(
                "Missing merged ranges:\n"
                + "\n".join(f"  - {value}" for value in missing_ranges)
            )

        if unexpected_ranges:
            details.append(
                "Unexpected merged ranges:\n"
                + "\n".join(f"  - {value}" for value in unexpected_ranges)
            )

        raise WorkbookValidationError(
            f"Invalid merged ranges.\n"
            f"File: {path}\n"
            f"Worksheet: {sheet.title}\n"
            + "\n".join(details)
            + "\nRemove unexpected merged ranges, "
            "then restore the missing merged ranges."
        )

def _normalize_formula(formula: str) -> str:
    """Normalize formula."""
    parts = []

    for token in Tokenizer(formula).items:
        value = token.value

        if (
            token.type == "OPERAND"
            and token.subtype == "RANGE"
            and "!" in value
        ):
            sheet_name, reference = value.rsplit("!", 1)

            if sheet_name.startswith("'") and sheet_name.endswith("'"):
                sheet_name = sheet_name[1:-1].replace("''", "'")

            value = f"{sheet_name}!{reference}"

        parts.append(value)

    return "=" + "".join(parts)

def _validate_labels_and_formulas(
    workbook: WorkbookType,
    reference: WorkbookType,
    path: Path,
    months_name: list[str]
) -> None:
    """Validate labels and formulas."""
    sheet_names = [*months_name, "Summary"]

    # Validate labels before formulas.
    for check_formulas in (False, True):
        for sheet_name in sheet_names:
            sheet = workbook[sheet_name]
            reference_sheet = reference[sheet_name]

            columns = _get_validation_columns(
                sheet=reference_sheet,
                months_name=months_name,
            )

            if sheet_name == "Summary":
                rows = range(4, 44)
            else:
                rows = [*range(4, 44), *range(46, 86)]

            for row in rows:
                for column in columns:
                    expected_cell = reference_sheet.cell(row=row, column=column)
                    expected_value = expected_cell.value

                    if not isinstance(expected_value, str):
                        continue

                    is_formula = expected_cell.data_type == "f"

                    if is_formula != check_formulas:
                        continue

                    actual_cell = sheet.cell(row=row, column=column)

                    if actual_cell.value == expected_value:
                        continue

                    if (
                        is_formula
                        and actual_cell.data_type == "f"
                        and _normalize_formula(actual_cell.value)
                        == _normalize_formula(expected_value)
                    ):
                        continue

                    description = (
                        "Invalid or missing formula"
                        if check_formulas else
                        "Invalid or missing label"
                    )

                    raise WorkbookValidationError(
                        f"{description}.\n"
                        f"File: {path}\n"
                        f"Worksheet: {sheet_name}\n"
                        f"Cell: {actual_cell.coordinate}\n"
                        f"Current value: {actual_cell.value!r}\n"
                        f"Expected value: {expected_value!r}\n"
                        "Restore the cell to the expected value shown."
                    )

def _validate_empty_cells(
    workbook: WorkbookType,
    reference: WorkbookType,
    path: Path,
    months_name: list[str]
) -> None:
    """Validate empty cells."""
    for sheet_name in [*months_name, "Summary"]:
        sheet = workbook[sheet_name]
        reference_sheet = reference[sheet_name]

        columns = _get_validation_columns(
                    sheet=reference_sheet,
                    months_name=months_name,
        )

        if sheet_name == "Summary":
            rows = range(4, 44)
        else:
            rows = [*range(4, 44), *range(46, 86)]

        for row in rows:
            reference_values = [
                reference_sheet.cell(row=row, column=column).value
                for column in columns
            ]

            if all(value is None for value in reference_values):
                columns = columns
            elif reference_values[0] is None:
                columns = (1,)
            else:
                continue

            for column in columns:
                cell = sheet.cell(row=row, column=column)

                if cell.value is not None:
                    raise WorkbookValidationError(
                        f"A cell that must remain empty contains a value.\n"
                        f"File: {path}\n"
                        f"Worksheet: {sheet_name}\n"
                        f"Cell: {cell.coordinate}\n"
                        f"Current value: {cell.value!r}\n"
                        "Clear the cell contents. "
                        "If the table has moved, restore "
                        "the original layout."
                    )


def _validate_numeric_cells(
    workbook: WorkbookType,
    reference: WorkbookType,
    path: Path,
    months_name: list[str]
) -> None:
    """Validate numeric cells."""
    for sheet_name in [*months_name, "Summary"]:
        sheet = workbook[sheet_name]
        reference_sheet = reference[sheet_name]

        columns = _get_validation_columns(
            sheet=reference_sheet,
            months_name=months_name,
        )

        if sheet_name == "Summary":
            rows = range(4, 44)
        else:
            rows = [*range(4, 41), *range(46, 86)]

        for row in rows:
            reference_values = [
                reference_sheet.cell(row=row, column=column).value
                for column in columns
            ]

            # Separator rows have already been validated.
            if all(value is None for value in reference_values):
                continue

            for column in columns:
                if column == 1:
                    continue

                reference_cell = reference_sheet.cell(
                    row=row,
                    column=column,
                )

                if isinstance(reference_cell, MergedCell):
                    continue

                expected_value = reference_cell.value

                # Fixed labels and formulas have already been validated.
                if isinstance(expected_value, str):
                    continue

                cell = sheet.cell(row=row, column=column)
                value = cell.value

                if value is None:
                    continue

                if type(value) in (int, float):
                    if isinstance(value, int) or isfinite(value):
                        continue

                raise WorkbookValidationError(
                    f"Invalid data cell content.\n"
                    f"File: {path}\n"
                    f"Worksheet: {sheet_name}\n"
                    f"Cell: {cell.coordinate}\n"
                    f"Current value: {value!r}\n"
                    "The cell must contain a number or remain empty. "
                    "Correct its contents."
                )

def _get_column_width(
    sheet: Worksheet,
    column: int,
) -> float:
    """Get column width."""
    for letter, dimension in sheet.column_dimensions.items():
        first_column = dimension.min or column_index_from_string(letter)
        last_column = dimension.max or first_column

        if first_column <= column <= last_column:
            return dimension.width

    if sheet.sheet_format.defaultColWidth is not None:
        return sheet.sheet_format.defaultColWidth

    return 13.0

def _validate_formatting_and_dimensions(
    workbook: WorkbookType,
    reference: WorkbookType,
    path: Path,
    months_name: list[str]
) -> None:
    """Warn about meaningful differences in styles, row heights and column widths."""
    style_attributes = {
        "font": "font",
        "fill": "fill",
        "border": "border",
        "alignment": "alignment",
        "number_format": "number format",
        "protection": "cell protection",
    }

    for reference_sheet in reference:
        sheet = workbook[reference_sheet.title]

        columns = _get_validation_columns(
            sheet=reference_sheet,
            months_name=months_name
        )

        for row in range(1, reference_sheet.max_row + 1):
            expected_height = reference_sheet.row_dimensions[row].height

            if expected_height is None:
                expected_height = reference_sheet.sheet_format.defaultRowHeight

            actual_height = sheet.row_dimensions[row].height

            if actual_height is None:
                actual_height = sheet.sheet_format.defaultRowHeight

            if not isclose(
                actual_height,
                expected_height,
                rel_tol=0,
                abs_tol=0.1,
            ):
                warnings.warn(
                    (
                        f"Invalid row height.\n"
                        f"File: {path}\n"
                        f"Worksheet: {sheet.title}\n"
                        f"Row: {row}\n"
                        f"Current height: {actual_height!r}\n"
                        f"Expected height: {expected_height!r}"
                    ),
                    category=WorkbookValidationWarning,
                )

        for column in columns:
            letter = get_column_letter(column)
            expected_width = _get_column_width(reference_sheet, column)
            actual_width = _get_column_width(sheet, column)

            if not isclose(
                actual_width,
                expected_width,
                rel_tol=0,
                abs_tol=0.1,
            ):
                warnings.warn(
                    (
                        f"Invalid column width.\n"
                        f"File: {path}\n"
                        f"Worksheet: {sheet.title}\n"
                        f"Column: {letter}\n"
                        f"Current width: {actual_width!r}\n"
                        f"Expected width: {expected_width!r}"
                    ),
                    category=WorkbookValidationWarning,
                )

        for row in range(1, reference_sheet.max_row + 1):
            for column in columns:
                expected_cell = reference_sheet.cell(row=row, column=column)

                actual_cell = sheet.cell(
                    row=expected_cell.row,
                    column=expected_cell.column,
                )

                for attribute, description in style_attributes.items():
                    expected_style = copy(getattr(expected_cell, attribute))
                    actual_style = copy(getattr(actual_cell, attribute))

                    if attribute == "alignment":
                        for alignment in (expected_style, actual_style):
                            if alignment.horizontal is None:
                                alignment.horizontal = "general"
                            if alignment.vertical is None:
                                alignment.vertical = "bottom"

                    elif attribute in ("font", "fill"):
                        if attribute == "font":
                            color_pairs = [
                                (expected_style.color, actual_style.color)
                            ]
                        else:
                            color_pairs = [
                                (expected_style.fgColor, actual_style.fgColor),
                                (expected_style.bgColor, actual_style.bgColor)
                            ]

                        for expected_color, actual_color in color_pairs:
                            if expected_color is None or actual_color is None:
                                continue

                            if isclose(
                                actual_color.tint,
                                expected_color.tint,
                                rel_tol=0,
                                abs_tol=0.00005
                            ):
                                actual_color.tint = expected_color.tint

                    elif attribute == "border":
                        for side_name in ("left", "right", "top", "bottom"):
                            for border in (expected_style, actual_style):
                                side = getattr(border, side_name)

                                if (
                                    side is not None
                                    and side.style is not None
                                    and side.color is None
                                ):
                                    side.color = Color(auto=True)

                    elif attribute == "number_format":
                        date_formats = {
                            "yyyy-mm-dd",
                            r"yyyy\-mm\-dd"
                        }

                        if (
                            expected_style in date_formats
                            and actual_style in date_formats
                        ):
                            actual_style = expected_style

                    if actual_style != expected_style:
                        warnings.warn(
                            (
                            f"Invalid cell formatting: {description}.\n"
                            f"File: {path}\n"
                            f"Worksheet: {sheet.title}\n"
                            f"Cell: {actual_cell.coordinate}\n"
                            "Restore the cell formatting "
                            "to match the report template."
                            ),
                            category=WorkbookValidationWarning,
                        )

def _year_workbook_path(directory: str | Path, year: int) -> Path:
    """Year workbook path."""
    return Path(directory) / f"Sales_Report_{year}.xlsx"

def load_year_workbook(
    directory: str | Path,
    year: int,
    months: set[int],
    template_sheet: Worksheet,
) -> tuple[WorkbookType | None, Path]:
    """Load and validate an annual report, or create it when it does not exist."""
    path = _year_workbook_path(directory, year)

    try:
        file_exists = path.exists()
    except OSError as exc:
        raise WorkbookFileAccessError(
            f"Cannot check annual report availability.\n"
            f"File: {path}\n"
            f"Details: {exc}"
        ) from exc

    if file_exists:
        try:
            workbook = load_workbook(path, data_only=False)
        except (InvalidFileException, BadZipFile, ParseError) as exc:
            raise WorkbookValidationError(
                f"The report cannot be opened as an Excel workbook.\n"
                f"File: {path}\n"
                f"Details: {exc}\n"
                "Check whether it can be opened in Excel. "
                "If so, save it again in .xlsx format."
            ) from exc
        except OSError as exc:
            raise WorkbookFileAccessError(
                f"Could not open or read the annual report.\n"
                f"File: {path}\n"
                f"Details: {exc}\n"
                "Check file availability and read permissions."
            ) from exc

        validation_completed = False
        reference = None
        months_name = [MONTH_NAMES[month - 1] for month in sorted(months)]  

        start = perf_counter()

        try:
            _validate_required_sheets(workbook, path, months_name)

            for month_name in months_name:
                _validate_month_headers(
                    sheet=workbook[month_name],
                    year=year,
                    month_name=month_name,
                    path=path,
                )

            _validate_other_sheet_headers(workbook, path)

            template = create_template_workbook() 

            try:
                reference = create_partial_year_workbook(
                    template_sheet=template["September"],
                    year=year,
                    months=months,
                )
            finally:
                template.close()

            _validate_merged_ranges(workbook, reference, path)
            _validate_labels_and_formulas(workbook, reference, path, months_name)
            _validate_empty_cells(workbook, reference, path, months_name)
            _validate_numeric_cells(workbook, reference, path, months_name)
            _validate_formatting_and_dimensions(workbook, reference, path, months_name)

            validation_completed = True

        finally:
            if reference is not None:
                reference.close()

            if not validation_completed:
                workbook.close()

            elapsed = perf_counter() - start
            print(f"Workbook validation duration: {elapsed:.2f} seconds")


    else:
        start = perf_counter()
        try:
            workbook = create_year_workbook(
                template_sheet=template_sheet,
                year=year,
            )
        finally:
            elapsed = perf_counter() - start
            print(f"Workbook creation duration: {elapsed:.2f} seconds") 

    return workbook, path

def save_workbook(
    path: str | Path,
    workbook: WorkbookType,
) -> None:
    """Save through a temporary file, then replace the destination workbook."""
    path = Path(path)
    temporary_path = None

    try:
        path.parent.mkdir(parents=True, exist_ok=True)

        with NamedTemporaryFile(
            dir=path.parent,
            prefix=f".{path.stem}_",
            suffix=".xlsx",
            delete=False,
        ) as temporary_file:
            temporary_path = Path(temporary_file.name)

        workbook.save(temporary_path)
        temporary_path.replace(path)

    except OSError as exc:
        raise WorkbookSaveError(
            f"Could not save the annual report.\n"
            f"File: {path}\n"
            f"Details: {exc}\n"
            "Check write permissions, available disk space "
            "and whether another application has locked the file."
        ) from exc

    finally:
        if temporary_path is not None:
            try:
                temporary_path.unlink(missing_ok=True)
            except OSError:
                pass