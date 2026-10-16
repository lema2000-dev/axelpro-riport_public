"""Sablon és éves riport betöltése, mentés ideiglenes fájlon keresztül."""

from pathlib import Path
from tempfile import NamedTemporaryFile
from time import perf_counter
from zipfile import BadZipFile
from xml.etree.ElementTree import ParseError
from datetime import date, datetime
from math import isfinite, isclose
from copy import copy
import warnings

from openpyxl import load_workbook
from openpyxl.workbook.workbook import Workbook as WorkbookType
from openpyxl.worksheet.worksheet import Worksheet
from openpyxl.utils import get_column_letter, column_index_from_string
from openpyxl.utils.exceptions import InvalidFileException
from openpyxl.cell.cell import MergedCell
from openpyxl.formula import Tokenizer
from openpyxl.styles import Color

from .excel_create_workbook import create_partial_year_workbook, create_year_workbook
from .excel_constants import (
    MONTH_NAMES,
    DAILY_EXCHANGE_RATES_SHEET,
    PROCESSED_DOCUMENTS_SHEET,
)
from .calendar_utils import get_month_weeks
from .template_creator import create_template_workbook

class WorkbookValidationError(ValueError):
    """Az éves riport fájla vagy tartalma hibás."""

class WorkbookValidationWarning(Warning):
    """Az éves riport fájla vagy tartalma figyelmeztetést érdemel."""

class WorkbookFileAccessError(OSError):
    """Az éves riport fájl elérése vagy olvasása sikertelen."""

class WorkbookSaveError(OSError):
    """Az éves riport fájl mentése sikertelen."""

def _get_validation_columns(
    sheet: Worksheet,
    months_name: list[str],
) -> list[int]:
    """Az ellenőrzendő oszlopok meghatározása."""
    if sheet.title == "Összesítő":
        return [
            1,
            *sorted(
                MONTH_NAMES.index(month_name) + 2
                for month_name in months_name
            ),
        ]

    return list(range(1, sheet.max_column + 1))

def _validate_required_sheets(
    workbook: WorkbookType,
    path: Path,
    months_name: list[str]
) -> None:
    """Az éves riport kötelező munkalapjainak ellenőrzése."""

    required_sheets = [
        *months_name,
        "Összesítő",
        "Euro",
        DAILY_EXCHANGE_RATES_SHEET,
        PROCESSED_DOCUMENTS_SHEET,
    ]

    missing_sheets = [
        name
        for name in required_sheets
        if name not in workbook.sheetnames
    ]

    if missing_sheets:
        missing_names = "\n".join(
            f"  - {name}" for name in missing_sheets
        )

        raise WorkbookValidationError(
            f"Hiányzó vagy átnevezett munkalapok.\n"
            f"Fájl: {path}\n"
            f"Elvárt munkalapnevek, amelyek nem találhatók:\n"
            f"{missing_names}\n"
            "Ellenőrizd a munkalapneveket. Az átnevezett lapokat "
            "nevezd vissza; a törölt lapokat állítsd helyre."
        )

def _validate_header_cell(
    sheet: Worksheet,
    row: int,
    column: int,
    expected_value: str | int | date | None,
    path: Path,
) -> None:
    """Egy fejléccella értékének ellenőrzése."""
    cell = sheet.cell(row=row, column=column)
    actual_value = cell.value

    if isinstance(expected_value, date) and isinstance(actual_value, datetime):
        actual_value = actual_value.date()

    if actual_value != expected_value:
        raise WorkbookValidationError(
            f"Hibás fejléc vagy elmozdult táblarész.\n"
            f"Fájl: {path}\n"
            f"Munkalap: {sheet.title}\n"
            f"Cella: {cell.coordinate}\n"
            f"Jelenlegi érték: {cell.value!r}\n"
            f"Elvárt érték: {expected_value!r}\n"
            "Állítsd helyre a megadott cella értékét. "
            "Ha sorokat vagy oszlopokat szúrtál be vagy töröltél, "
            "a táblázat eredeti elrendezését is állítsd helyre."
        )

def _validate_month_headers(
    sheet: Worksheet,
    year: int,
    month_name: str,
    path: Path,
) -> None:
    """A havi munkalap heti blokkjainak és havi fejléceinek ellenőrzése."""
    month = MONTH_NAMES.index(month_name) + 1
    weekday_names = (
        "Hétfő",
        "Kedd",
        "Szerda",
        "Csütörtök",
        "Péntek",
        "Szombat",
    )

    expected_cells = {
        (1, 1): "Dátum:",
        (3, 1): "Pénznem:",
        (43, 1): "Dátum:",
        (43, 2): "Havi záró",
        (44, 2): MONTH_NAMES[month - 1],
        (45, 1): "Pénznem:",
        (45, 2): "Forint",
    }

    next_column = 2

    for week in get_month_weeks(year, month):
        days = [day for day in week if day.weekday() < 6]

        if not days:
            continue

        for day in days:
            expected_cells[(1, next_column)] = weekday_names[day.weekday()]
            expected_cells[(2, next_column)] = day
            expected_cells[(3, next_column)] = "Forint"
            next_column += 2

        expected_cells[(1, next_column)] = "Heti záró"
        expected_cells[(2, next_column)] = days[0].isocalendar().week
        expected_cells[(3, next_column)] = "Forint"
        next_column += 2

    # A fejlécek üres celláiban sem lehet eltérő érték.
    for row in (1, 2, 3):
        for column in range(1, next_column):
            expected_cells.setdefault((row, column), None)

    for row in (43, 44, 45):
        for column in range(1, 4):
            expected_cells.setdefault((row, column), None)

    for (row, column), expected_value in expected_cells.items():
        _validate_header_cell(
            sheet=sheet,
            row=row,
            column=column,
            expected_value=expected_value,
            path=path,
        )

def _validate_other_sheet_headers(
    workbook: WorkbookType,
    path: Path,
) -> None:
    """Az összesítő és a nyilvántartások fejléceinek ellenőrzése."""
    summary = workbook["Összesítő"]

    for row, expected_value in (
        (1, "Dátum:"),
        (2, None),
        (3, "Pénznem:"),
    ):
        _validate_header_cell(
            sheet=summary,
            row=row,
            column=1,
            expected_value=expected_value,
            path=path,
        )

    for column, month_name in enumerate(MONTH_NAMES, start=2):
        for row, expected_value in (
            (1, "Havi záró"),
            (2, month_name),
            (3, "Forint")
        ):
            _validate_header_cell(
                sheet=summary,
                row=row,
                column=column,
                expected_value=expected_value,
                path=path,
            )

    headers = {
        "Euro": (
            "Hónap",
            "EUR/HUF árfolyam",
            "Árfolyam dátuma",
        ),
        DAILY_EXCHANGE_RATES_SHEET: (
            "Alkalmazás dátuma",
            "Árfolyam dátuma",
            "EUR/HUF árfolyam",
        ),
        PROCESSED_DOCUMENTS_SHEET: (
            "Év",
            "Bizonylatszám",
        ),
    }

    for sheet_name, column_headers in headers.items():
        for column, expected_value in enumerate(column_headers, start=1):
            _validate_header_cell(
                sheet=workbook[sheet_name],
                row=1,
                column=column,
                expected_value=expected_value,
                path=path,
            )

    euro = workbook["Euro"]

    for row, month_name in enumerate(MONTH_NAMES, start=2):
        _validate_header_cell(
            sheet=euro,
            row=row,
            column=1,
            expected_value=month_name,
            path=path,
        )

def _validate_merged_ranges(
    workbook: WorkbookType,
    reference: WorkbookType,
    path: Path,   
) -> None:
    """A cellaösszevonások ellenőrzése az éves mintához képest."""
    for reference_sheet in reference:
        sheet = workbook[reference_sheet.title]

        expected_ranges = {
            str(merged_range)
            for merged_range in reference_sheet.merged_cells.ranges
        }

        actual_ranges = {
            str(merged_range)
            for merged_range in sheet.merged_cells.ranges
        }

        missing_ranges = sorted(expected_ranges - actual_ranges)
        unexpected_ranges = sorted(actual_ranges - expected_ranges)

        if not missing_ranges and not unexpected_ranges:
            continue

        details = []

        if missing_ranges:
            details.append(
                "Hiányzó cellaösszevonások:\n"
                + "\n".join(f"  - {value}" for value in missing_ranges)
            )

        if unexpected_ranges:
            details.append(
                "Nem megfelelő cellaösszevonások:\n"
                + "\n".join(f"  - {value}" for value in unexpected_ranges)
            )

        raise WorkbookValidationError(
            f"Hibás cellaösszevonások.\n"
            f"Fájl: {path}\n"
            f"Munkalap: {sheet.title}\n"
            + "\n".join(details)
            + "\nSzüntesd meg a nem megfelelő cellaösszevonásokat, "
            "majd állítsd helyre a hiányzó összevonásokat."
        )

def _normalize_formula(formula: str) -> str:
    """Munkalaphivatkozások idézésének egységesítése összehasonlításhoz."""
    parts = []

    for token in Tokenizer(formula).items:
        value = token.value

        if (
            token.type == "OPERAND"
            and token.subtype == "RANGE"
            and "!" in value
        ):
            sheet_name, reference = value.rsplit("!", 1)

            if sheet_name.startswith("'") and sheet_name.endswith("'"):
                sheet_name = sheet_name[1:-1].replace("''", "'")

            value = f"{sheet_name}!{reference}"

        parts.append(value)

    return "=" + "".join(parts)

def _validate_labels_and_formulas(
    workbook: WorkbookType,
    reference: WorkbookType,
    path: Path,
    months_name: list[str]
) -> None:
    """A rögzített feliratok és képletek ellenőrzése az éves mintához képest."""
    sheet_names = [*months_name, "Összesítő"]

    # Először a feliratokat, majd a képleteket ellenőrizzük.
    for check_formulas in (False, True):
        for sheet_name in sheet_names:
            sheet = workbook[sheet_name]
            reference_sheet = reference[sheet_name]

            columns = _get_validation_columns(
                sheet=reference_sheet,
                months_name=months_name,
            )

            if sheet_name == "Összesítő":
                rows = range(4, 44)
            else:
                rows = [*range(4, 44), *range(46, 86)]

            for row in rows:
                for column in columns:
                    expected_cell = reference_sheet.cell(row=row, column=column)
                    expected_value = expected_cell.value

                    if not isinstance(expected_value, str):
                        continue

                    is_formula = expected_cell.data_type == "f"

                    if is_formula != check_formulas:
                        continue

                    actual_cell = sheet.cell(row=row, column=column)

                    if actual_cell.value == expected_value:
                        continue

                    if (
                        is_formula
                        and actual_cell.data_type == "f"
                        and _normalize_formula(actual_cell.value)
                        == _normalize_formula(expected_value)
                    ):
                        continue

                    description = (
                        "Hibás vagy hiányzó képlet"
                        if check_formulas else
                        "Hibás vagy hiányzó felirat"
                    )

                    raise WorkbookValidationError(
                        f"{description}.\n"
                        f"Fájl: {path}\n"
                        f"Munkalap: {sheet_name}\n"
                        f"Cella: {actual_cell.coordinate}\n"
                        f"Jelenlegi érték: {actual_cell.value!r}\n"
                        f"Elvárt érték: {expected_value!r}\n"
                        "Állítsd helyre a cellát a megadott elvárt értékkel."
                    )

def _validate_empty_cells(
    workbook: WorkbookType,
    reference: WorkbookType,
    path: Path,
    months_name: list[str]
) -> None:
    """Az elválasztó sorok, és az üres felirathelyek ellenőrzése."""
    for sheet_name in [*months_name, "Összesítő"]:
        sheet = workbook[sheet_name]
        reference_sheet = reference[sheet_name]

        columns = _get_validation_columns(
                    sheet=reference_sheet,
                    months_name=months_name,
        )

        if sheet_name == "Összesítő":
            rows = range(4, 44)
        else:
            rows = [*range(4, 44), *range(46, 86)]

        for row in rows:
            reference_values = [
                reference_sheet.cell(row=row, column=column).value
                for column in columns
            ]

            if all(value is None for value in reference_values):
                columns = columns
            elif reference_values[0] is None:
                columns = (1,)
            else:
                continue

            for column in columns:
                cell = sheet.cell(row=row, column=column)

                if cell.value is not None:
                    raise WorkbookValidationError(
                        f"Érték került egy üresen tartandó cellába.\n"
                        f"Fájl: {path}\n"
                        f"Munkalap: {sheet_name}\n"
                        f"Cella: {cell.coordinate}\n"
                        f"Jelenlegi érték: {cell.value!r}\n"
                        "Töröld a cella tartalmát. "
                        "Ha a táblázat elmozdult, állítsd helyre "
                        "az eredeti elrendezést."
                    )


def _validate_numeric_cells(
    workbook: WorkbookType,
    reference: WorkbookType,
    path: Path,
    months_name: list[str]
) -> None:
    """A kitölthető adatcellák számértékeinek ellenőrzése."""
    for sheet_name in [*months_name, "Összesítő"]:
        sheet = workbook[sheet_name]
        reference_sheet = reference[sheet_name]

        columns = _get_validation_columns(
            sheet=reference_sheet,
            months_name=months_name,
        )

        if sheet_name == "Összesítő":
            rows = range(4, 44)
        else:
            rows = [*range(4, 41), *range(46, 86)]

        for row in rows:
            reference_values = [
                reference_sheet.cell(row=row, column=column).value
                for column in columns
            ]

            # Az elválasztó sorokat már ellenőriztük.
            if all(value is None for value in reference_values):
                continue

            for column in columns:
                if column == 1:
                    continue

                reference_cell = reference_sheet.cell(
                    row=row,
                    column=column,
                )

                if isinstance(reference_cell, MergedCell):
                    continue

                expected_value = reference_cell.value

                # A rögzített feliratokat és képleteket már ellenőriztük.
                if isinstance(expected_value, str):
                    continue

                cell = sheet.cell(row=row, column=column)
                value = cell.value

                if value is None:
                    continue

                if type(value) in (int, float):
                    if isinstance(value, int) or isfinite(value):
                        continue

                raise WorkbookValidationError(
                    f"Hibás adatcellatartalom.\n"
                    f"Fájl: {path}\n"
                    f"Munkalap: {sheet_name}\n"
                    f"Cella: {cell.coordinate}\n"
                    f"Jelenlegi érték: {value!r}\n"
                    "A cella csak számot tartalmazhat, vagy üresen maradhat. "
                    "Javítsd a tartalmát."
                )

def _get_column_width(
    sheet: Worksheet,
    column: int,
) -> float:
    """Oszlopszélesség kiolvasása a közös oszloptartományokból is."""
    for letter, dimension in sheet.column_dimensions.items():
        first_column = dimension.min or column_index_from_string(letter)
        last_column = dimension.max or first_column

        if first_column <= column <= last_column:
            return dimension.width

    if sheet.sheet_format.defaultColWidth is not None:
        return sheet.sheet_format.defaultColWidth

    return 13.0

def _validate_formatting_and_dimensions(
    workbook: WorkbookType,
    reference: WorkbookType,
    path: Path,
    months_name: list[str]
) -> None:
    """A riport cellaformázásainak és méreteinek ellenőrzése."""
    style_attributes = {
        "font": "betűformázás",
        "fill": "kitöltés",
        "border": "szegély",
        "alignment": "igazítás",
        "number_format": "számformátum",
        "protection": "cellavédelem",
    }

    for reference_sheet in reference:
        sheet = workbook[reference_sheet.title]

        columns = _get_validation_columns(
            sheet=reference_sheet,
            months_name=months_name
        )

        for row in range(1, reference_sheet.max_row + 1):
            expected_height = reference_sheet.row_dimensions[row].height

            if expected_height is None:
                expected_height = reference_sheet.sheet_format.defaultRowHeight

            actual_height = sheet.row_dimensions[row].height

            if actual_height is None:
                actual_height = sheet.sheet_format.defaultRowHeight

            if not isclose(
                actual_height,
                expected_height,
                rel_tol=0,
                abs_tol=0.1,
            ):
                warnings.warn(
                    (
                        f"Hibás sormagasság.\n"
                        f"Fájl: {path}\n"
                        f"Munkalap: {sheet.title}\n"
                        f"Sor: {row}\n"
                        f"Jelenlegi magasság: {actual_height!r}\n"
                        f"Elvárt magasság: {expected_height!r}"
                    ),
                    category=WorkbookValidationWarning,
                )

        for column in columns:
            letter = get_column_letter(column)
            expected_width = _get_column_width(reference_sheet, column)
            actual_width = _get_column_width(sheet, column)

            if not isclose(
                actual_width,
                expected_width,
                rel_tol=0,
                abs_tol=0.1,
            ):
                warnings.warn(
                    (
                        f"Hibás oszlopszélesség.\n"
                        f"Fájl: {path}\n"
                        f"Munkalap: {sheet.title}\n"
                        f"Oszlop: {letter}\n"
                        f"Jelenlegi szélesség: {actual_width!r}\n"
                        f"Elvárt szélesség: {expected_width!r}"
                    ),
                    category=WorkbookValidationWarning,
                )

        for row in range(1, reference_sheet.max_row + 1):
            for column in columns:
                expected_cell = reference_sheet.cell(row=row, column=column)

                actual_cell = sheet.cell(
                    row=expected_cell.row,
                    column=expected_cell.column,
                )

                for attribute, description in style_attributes.items():
                    expected_style = copy(getattr(expected_cell, attribute))
                    actual_style = copy(getattr(actual_cell, attribute))

                    if attribute == "alignment":
                        for alignment in (expected_style, actual_style):
                            if alignment.horizontal is None:
                                alignment.horizontal = "general"
                            if alignment.vertical is None:
                                alignment.vertical = "bottom"

                    elif attribute in ("font", "fill"):
                        if attribute == "font":
                            color_pairs = [
                                (expected_style.color, actual_style.color)
                            ]
                        else:
                            color_pairs = [
                                (expected_style.fgColor, actual_style.fgColor),
                                (expected_style.bgColor, actual_style.bgColor)
                            ]

                        for expected_color, actual_color in color_pairs:
                            if expected_color is None or actual_color is None:
                                continue

                            if isclose(
                                actual_color.tint,
                                expected_color.tint,
                                rel_tol=0,
                                abs_tol=0.00005
                            ):
                                actual_color.tint = expected_color.tint

                    elif attribute == "border":
                        for side_name in ("left", "right", "top", "bottom"):
                            for border in (expected_style, actual_style):
                                side = getattr(border, side_name)

                                if (
                                    side is not None
                                    and side.style is not None
                                    and side.color is None
                                ):
                                    side.color = Color(auto=True)

                    elif attribute == "number_format":
                        date_formats = {
                            "yyyy-mm-dd",
                            r"yyyy\-mm\-dd"
                        }

                        if (
                            expected_style in date_formats
                            and actual_style in date_formats
                        ):
                            actual_style = expected_style

                    if actual_style != expected_style:
                        warnings.warn(
                            (
                            f"Hibás cellaformázás: {description}.\n"
                            f"Fájl: {path}\n"
                            f"Munkalap: {sheet.title}\n"
                            f"Cella: {actual_cell.coordinate}\n"
                            "Állítsd helyre a cella formázását "
                            "az eredeti riportmintának megfelelően."
                            ),
                            category=WorkbookValidationWarning,
                        )

def _year_workbook_path(directory: str | Path, year: int) -> Path:
    """Az éves riport célútvonalának összeállítása."""
    return Path(directory) / f"Napi_értékesítési_riport_{year}.xlsx"

def load_year_workbook(
    directory: str | Path,
    year: int,
    months: set[int],
    template_sheet: Worksheet,
) -> tuple[WorkbookType | None, Path]:
    """Éves riport betöltése, szükség esetén létrehozással."""
    path = _year_workbook_path(directory, year)

    try:
        file_exists = path.exists()
    except OSError as exc:
        raise WorkbookFileAccessError(
            f"Az éves riport elérhetősége nem ellenőrizhető.\n"
            f"Fájl: {path}\n"
            f"Részletek: {exc}"
        ) from exc

    if file_exists:
        try:
            workbook = load_workbook(path, data_only=False)
        except (InvalidFileException, BadZipFile, ParseError) as exc:
            raise WorkbookValidationError(
                f"A riport nem nyitható meg Excel-munkafüzetként.\n"
                f"Fájl: {path}\n"
                f"Részletek: {exc}\n"
                "Ellenőrizd, hogy Excelben megnyitható-e. "
                "Ha igen, mentsd újra .xlsx formátumban."
            ) from exc
        except OSError as exc:
            raise WorkbookFileAccessError(
                f"Az éves riport megnyitása vagy olvasása nem sikerült.\n"
                f"Fájl: {path}\n"
                f"Részletek: {exc}\n"
                "Ellenőrizd a fájl elérhetőségét és olvasási jogosultságát."
            ) from exc

        validation_completed = False
        reference = None
        months_name = [MONTH_NAMES[month - 1] for month in sorted(months)]  

        start = perf_counter()

        try:
            _validate_required_sheets(workbook, path, months_name)

            for month_name in months_name:
                _validate_month_headers(
                    sheet=workbook[month_name],
                    year=year,
                    month_name=month_name,
                    path=path,
                )

            _validate_other_sheet_headers(workbook, path)

            template = create_template_workbook() 

            try:
                reference = create_partial_year_workbook(
                    template_sheet=template["Szeptember"],
                    year=year,
                    months=months,
                )
            finally:
                template.close()

            _validate_merged_ranges(workbook, reference, path)
            _validate_labels_and_formulas(workbook, reference, path, months_name)
            _validate_empty_cells(workbook, reference, path, months_name)
            _validate_numeric_cells(workbook, reference, path, months_name)
            _validate_formatting_and_dimensions(workbook, reference, path, months_name)

            validation_completed = True

        finally:
            if reference is not None:
                reference.close()

            if not validation_completed:
                workbook.close()

            elapsed = perf_counter() - start
            print(f"Az munkafüzet ellenőrzésének futási ideje: {elapsed:.2f} másodperc")


    else:
        start = perf_counter()
        try:
            workbook = create_year_workbook(
                template_sheet=template_sheet,
                year=year,
            )
        finally:
            elapsed = perf_counter() - start
            print(f"Az új munkafüzet létrehozásának futási ideje: {elapsed:.2f} másodperc") 

    return workbook, path

def save_workbook(
    path: str | Path,
    workbook: WorkbookType,
) -> None:
    """Mentés ideiglenes fájlba, majd a célfájl cseréje."""
    path = Path(path)
    temporary_path = None

    try:
        path.parent.mkdir(parents=True, exist_ok=True)

        with NamedTemporaryFile(
            dir=path.parent,
            prefix=f".{path.stem}_",
            suffix=".xlsx",
            delete=False,
        ) as temporary_file:
            temporary_path = Path(temporary_file.name)

        workbook.save(temporary_path)
        temporary_path.replace(path)

    except OSError as exc:
        raise WorkbookSaveError(
            f"Az éves riport mentése nem sikerült.\n"
            f"Fájl: {path}\n"
            f"Részletek: {exc}\n"
            "Ellenőrizd az írási jogosultságot, a szabad lemezterületet "
            "és azt, hogy a fájl nincs-e zárolva egy másik programban."
        ) from exc

    finally:
        if temporary_path is not None:
            try:
                temporary_path.unlink(missing_ok=True)
            except OSError:
                pass