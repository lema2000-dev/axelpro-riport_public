"""A már mentett évek megőrzése és az erőforrások lezárása."""

from datetime import date
from decimal import Decimal
from unittest.mock import Mock
import pytest
from openpyxl import Workbook, load_workbook
from axelpro_riport import import_service as service, settings
from axelpro_riport.excel_files import (
    WorkbookSaveError,
    WorkbookFileAccessError,
    WorkbookValidationError,
)
from axelpro_riport.excel_constants import (
    PROCESSED_DOCUMENTS_SHEET,
    DAILY_EXCHANGE_RATES_SHEET,
)
from axelpro_riport.models import DailyExchangeRate


def small_report(year):
    wb = Workbook()
    wb.active.title = "Január"
    wb.active["B2"] = date(year, 1, 5)
    wb.create_sheet(PROCESSED_DOCUMENTS_SHEET).append(["Év", "Bizonylatszám"])
    wb.create_sheet(DAILY_EXCHANGE_RATES_SHEET).append(
        ["Application", "Source", "Rate"]
    )
    wb.create_sheet("Euro")
    wb.close = Mock(wraps=wb.close)
    return wb


@pytest.mark.parametrize("stage", ["load_access", "load_validation", "save"])
def test_later_failure_reports_saved_year_and_closes_workbooks(
    tmp_path, settings_location, write_csv, movement, monkeypatch, stage
):
    settings.create_settings(str(tmp_path), [])
    csv = write_csv(
        tmp_path / "input.csv",
        [movement(day=date(2025, 1, 5)), movement(day=date(2026, 1, 5))],
    )
    today = date.today()
    monkeypatch.setattr(
        service,
        "get_exchange_rates",
        lambda **kwargs: {today: DailyExchangeRate(today, today, Decimal(400))},
    )
    # Ez a teszt az éves írási szakaszt izolálja; a bizonylatolvasónak külön tesztje van.
    monkeypatch.setattr(
        service, "_read_processed_documents", lambda **kwargs: set()
    )
    template = Workbook()
    template.active.title = "Szeptember"
    template.close = Mock(wraps=template.close)
    monkeypatch.setattr(service, "create_template_workbook", lambda: template)
    reports = []
    error_type = {
        "load_access": WorkbookFileAccessError,
        "load_validation": WorkbookValidationError,
        "save": WorkbookSaveError,
    }[stage]
    failure = error_type("second year failed")

    def load(directory, year, **kwargs):
        if year == 2026 and stage != "save":
            raise failure
        wb = small_report(year)
        reports.append(wb)
        return wb, directory_path / f"Napi_értékesítési_riport_{year}.xlsx"

    directory_path = tmp_path
    monkeypatch.setattr(service, "load_year_workbook", load)
    monkeypatch.setattr(service, "_confirm_import_preview", lambda **kwargs: True)
    original_save = service.save_workbook

    def save(path, workbook):
        if "2026" in path.name and stage == "save":
            raise failure
        original_save(path, workbook)

    monkeypatch.setattr(service, "save_workbook", save)
    with pytest.raises(error_type) as error:
        service.import_csv(csv)
    if stage == "save":
        assert "Már sikeresen elmentett évek: 2025" in str(error.value)
        assert "2026" in str(error.value)
        assert error.value.__cause__ is failure
        assert (tmp_path / "Napi_értékesítési_riport_2025.xlsx").exists()
    else:
        # All years are loaded before any processing or saving starts.
        assert error.value is failure
        assert not (tmp_path / "Napi_értékesítési_riport_2025.xlsx").exists()
    assert not (tmp_path / "Napi_értékesítési_riport_2026.xlsx").exists()
    template.close.assert_called_once()
    for report in reports:
        report.close.assert_called_once()


def test_document_reader_reads_shared_reports_without_closing(tmp_path):
    wb = small_report(2026)
    wb[PROCESSED_DOCUMENTS_SHEET].append([2026, "OLD"])
    documents = service._read_processed_documents({2026: (wb, tmp_path / "report.xlsx")})
    assert documents == {(2026, "OLD")}
    wb.close.assert_not_called()
    assert service._read_processed_documents({}) == set()


def test_document_reader_does_not_close_on_failure(tmp_path):
    wb = Workbook()
    wb.close = Mock(wraps=wb.close)
    with pytest.raises(KeyError):
        service._read_processed_documents({2026: (wb, tmp_path / "report.xlsx")})
    wb.close.assert_not_called()


@pytest.mark.parametrize("stage", ["rates", "documents", "processor", "no_new", "reject"])
def test_import_closes_all_shared_workbooks_on_early_exit_or_error(
    tmp_path, settings_location, movement, monkeypatch, capsys, stage
):
    settings.create_settings(str(tmp_path), [])
    reports = {year: (small_report(year), tmp_path / f"{year}.xlsx") for year in (2025, 2026)}
    monkeypatch.setattr(service, "read_movements", lambda **kwargs: [movement()])
    monkeypatch.setattr(service, "_loads_workbooks", lambda **kwargs: reports)
    failure = ValueError("processing failed")
    def fail(**kwargs):
        raise failure
    monkeypatch.setattr(service, "get_exchange_rates", fail if stage == "rates" else lambda **kwargs: {})
    monkeypatch.setattr(service, "_read_processed_documents", fail if stage == "documents" else lambda **kwargs: set())
    from axelpro_riport.models import ImportResult
    result = ImportResult([], [], set())
    monkeypatch.setattr(service, "process_movements", fail if stage == "processor" else lambda **kwargs: None if stage == "no_new" else result)
    monkeypatch.setattr(service, "write_import_result", lambda **kwargs: None)
    monkeypatch.setattr(service, "write_exchange_rates", lambda **kwargs: None)
    monkeypatch.setattr(service, "_confirm_import_preview", lambda **kwargs: False)
    save = Mock()
    monkeypatch.setattr(service, "save_workbook", save)
    if stage in ("rates", "documents", "processor"):
        with pytest.raises(ValueError) as error:
            service.import_csv(tmp_path / "input.csv")
        assert error.value is failure
    else:
        service.import_csv(tmp_path / "input.csv")
    save.assert_not_called()
    for wb, _ in reports.values():
        wb.close.assert_called_once()


def test_loads_each_year_once_with_one_shared_template(tmp_path, movement, monkeypatch):
    template = Workbook()
    template.active.title = "Szeptember"
    template.close = Mock(wraps=template.close)
    create = Mock(return_value=template)
    monkeypatch.setattr(service, "create_template_workbook", create)
    calls = []
    def load(**kwargs):
        calls.append(kwargs)
        return small_report(kwargs["year"]), tmp_path / f'{kwargs["year"]}.xlsx'
    monkeypatch.setattr(service, "load_year_workbook", load)
    rows = [movement(day=date(2026, 1, 5)), movement(day=date(2025, 1, 6)), movement(day=date(2026, 2, 2))]
    reports = service._loads_workbooks(rows, tmp_path)
    assert list(reports) == [2025, 2026]
    assert [call["year"] for call in calls] == [2025, 2026]
    assert all(call["template_sheet"] is template["Szeptember"] for call in calls)
    assert all("create_if_missing" not in call for call in calls)
    assert [call["months"] for call in calls] == [{1}, {1, 2}]
    create.assert_called_once()
    template.close.assert_called_once()
    for wb, _ in reports.values():
        wb.close.assert_not_called()
