"""Test partial workbook validation."""
from io import BytesIO
import warnings
import pytest
from openpyxl import load_workbook
from openpyxl.styles import Font
from sales_report import excel_files as f
from sales_report.excel_create_workbook import create_partial_year_workbook
from sales_report.template_creator import create_template_workbook
from sales_report.excel_constants import MONTH_NAMES


@pytest.fixture(scope="module")
def partial_bytes():
    template = create_template_workbook()
    try:
        reference = create_partial_year_workbook(template["September"], 2026, [1, 9])
        try:
            stream = BytesIO()
            reference.save(stream)
            return stream.getvalue()
        finally:
            reference.close()
    finally:
        template.close()


@pytest.fixture
def partial_reference(partial_bytes):
    reference = load_workbook(BytesIO(partial_bytes))
    yield reference
    reference.close()


def test_partial_reference_contains_only_requested_months(partial_reference, reference_report):
    assert [name for name in MONTH_NAMES if name in partial_reference.sheetnames] == ["January", "September"]
    for column in (2, 10):
        for row in range(1, 44):
            assert partial_reference["Summary"].cell(row, column).value == reference_report["Summary"].cell(row, column).value
    assert partial_reference["Summary"]["C4"].value is None
    for name in ("January", "September"):
        assert partial_reference[name].freeze_panes == "D4"
        assert partial_reference[name]["B85"].value is None
    assert partial_reference["Exchange_Rates"].max_row == 13
    assert partial_reference["_daily_exchange_rates"].max_row == 1
    assert partial_reference["_imported_documents"].max_row == 1


def test_summary_validation_columns_keep_original_month_positions(partial_reference):
    assert f._get_validation_columns(partial_reference["Summary"], ["September", "January"]) == [1, 2, 10]


def test_partial_reference_accepts_full_report_and_summary_gap_formulas(year_report, partial_reference, tmp_path):
    names = ["January", "September"]
    path = tmp_path / "report.xlsx"
    f._validate_merged_ranges(year_report, partial_reference, path)
    f._validate_labels_and_formulas(year_report, partial_reference, path, names)
    f._validate_empty_cells(year_report, partial_reference, path, names)
    f._validate_numeric_cells(year_report, partial_reference, path, names)
    with warnings.catch_warnings(record=True) as collected:
        warnings.simplefilter("always", f.WorkbookValidationWarning)
        f._validate_formatting_and_dimensions(year_report, partial_reference, path, names)
    assert not [w for w in collected if issubclass(w.category, f.WorkbookValidationWarning)]


def test_unselected_month_values_and_styles_are_outside_validation_scope(year_report, partial_reference, tmp_path):
    year_report["February"]["J8"] = "invalid but not selected"
    year_report["Summary"]["C4"] = "unselected column"
    year_report["Summary"]["C4"].font = Font(name="Different", size=30)
    year_report["Summary"].column_dimensions["C"].width = 99
    names = ["January", "September"]
    path = tmp_path / "report.xlsx"
    for validate in (f._validate_labels_and_formulas, f._validate_empty_cells, f._validate_numeric_cells):
        validate(year_report, partial_reference, path, names)
    with warnings.catch_warnings(record=True) as collected:
        warnings.simplefilter("always", f.WorkbookValidationWarning)
        f._validate_formatting_and_dimensions(year_report, partial_reference, path, names)
    assert not [w for w in collected if issubclass(w.category, f.WorkbookValidationWarning)]


def test_selected_summary_formula_is_still_checked(year_report, partial_reference, tmp_path):
    summary = partial_reference["Summary"]
    expected = next(summary.cell(row, 10) for row in range(4, 44) if summary.cell(row, 10).data_type == "f")
    year_report["Summary"][expected.coordinate] = "=1+2"
    with pytest.raises(f.WorkbookValidationError, match=expected.coordinate):
        f._validate_labels_and_formulas(year_report, partial_reference, tmp_path / "report.xlsx", ["January", "September"])


def test_required_sheets_check_selected_months_and_shared_sheets(year_report, tmp_path):
    del year_report["February"]
    f._validate_required_sheets(year_report, tmp_path / "report.xlsx", ["January", "September"])
    del year_report["September"]
    with pytest.raises(f.WorkbookValidationError, match="September"):
        f._validate_required_sheets(year_report, tmp_path / "report.xlsx", ["January", "September"])


def test_selected_month_formatting_still_warns(year_report, partial_reference, tmp_path):
    year_report["September"]["B8"].font = Font(name="Different", size=30)
    with pytest.warns(f.WorkbookValidationWarning, match="B8"):
        f._validate_formatting_and_dimensions(year_report, partial_reference, tmp_path / "report.xlsx", ["January", "September"])


def test_loading_builds_reference_for_selected_months_only(tmp_path, year_report_bytes, template_sheet, monkeypatch):
    path = tmp_path / "Sales_Report_2026.xlsx"
    path.write_bytes(year_report_bytes)
    original = f.create_partial_year_workbook
    calls = []
    def create(**kwargs):
        calls.append(kwargs["months"])
        return original(**kwargs)
    monkeypatch.setattr(f, "create_partial_year_workbook", create)
    loaded, result_path = f.load_year_workbook(tmp_path, 2026, {9}, template_sheet)
    try:
        assert calls == [{9}]
        assert result_path == path
        assert all(name in loaded.sheetnames for name in MONTH_NAMES)
    finally:
        loaded.close()
