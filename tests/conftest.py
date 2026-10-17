"""Conftest."""

import csv
import sys
from pathlib import Path
from datetime import date
from decimal import Decimal
import pytest
from openpyxl import Workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "src"))
from sales_report.models import Movement


@pytest.fixture
def movement():
    def make(
        day=date(2026, 1, 5),
        document="SZ-1",
        group="New parts",
        amount="100",
        currency="HUF",
        transaction="Invoice",
        item="PART",
    ):
        return Movement(
            day, transaction, document, group, item, Decimal(amount), currency
        )

    return make


@pytest.fixture
def write_csv():
    def write(path, rows, *, partners=None):
        columns = [
            "Movement timestamp",
            "Transaction type",
            "Document number",
            "Partner",
            "Group",
            "Item code",
            "Net value",
            "Currency",
        ]
        with path.open("w", encoding="utf-16", newline="") as stream:
            writer = csv.writer(stream, delimiter="\t")
            writer.writerow(columns)
            for index, m in enumerate(rows):
                writer.writerow(
                    [
                        m.movement_date.strftime("%Y. %m. %d. 12:00:00"),
                        m.transaction_type,
                        m.document_number,
                        "Example customer" if partners is None else partners[index],
                        m.group,
                        m.item_code,
                        str(m.net_value).replace(".", ","),
                        m.currency,
                    ]
                )
        return path

    return write


@pytest.fixture
def template_sheet():
    sheet = Workbook().active
    sheet.title = "September"
    for row in range(1, 86):
        sheet.cell(row, 1, f"Row {row}")
    for col in range(2, 14, 2):
        sheet.cell(2, col, date(2026, 9, 14 + (col - 2) // 2))
        sheet.cell(8, col, 999)
        sheet.cell(
            12, col, f"={sheet.cell(8, col).coordinate}+{sheet.cell(9, col).coordinate}"
        )
        sheet.merge_cells(start_row=8, end_row=8, start_column=col, end_column=col + 1)
    sheet["N12"] = "=B12+D12+F12+H12+J12+L12"
    sheet["B82"] = 999
    sheet["B49"] = "=B46+B47+B48"
    sheet.row_dimensions[8].height = 25
    sheet.column_dimensions["B"].width = 20
    return sheet


@pytest.fixture
def settings_location(monkeypatch, tmp_path):
    from sales_report import settings, console

    path = tmp_path / "config" / "settings.json"
    monkeypatch.setattr(settings, "settings_path", lambda: path)
    monkeypatch.setattr(console, "settings_path", lambda: path)
    return path


@pytest.fixture(scope="session")
def year_report_bytes():
    """Year report bytes."""
    from io import BytesIO
    from sales_report.template_creator import create_template_workbook
    from sales_report.excel_create_workbook import create_year_workbook

    template = create_template_workbook()
    try:
        workbook = create_year_workbook(template["September"], 2026)
        stream = BytesIO()
        workbook.save(stream)
        workbook.close()
        return stream.getvalue()
    finally:
        template.close()


@pytest.fixture
def year_report(year_report_bytes):
    from io import BytesIO
    from openpyxl import load_workbook

    workbook = load_workbook(BytesIO(year_report_bytes))
    yield workbook
    workbook.close()


@pytest.fixture
def reference_report(year_report_bytes):
    from io import BytesIO
    from openpyxl import load_workbook

    workbook = load_workbook(BytesIO(year_report_bytes))
    yield workbook
    workbook.close()
