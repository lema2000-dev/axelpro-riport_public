"""Valódi fájlok és determinisztikusan előidézett I/O hibák."""

from pathlib import Path
from xml.etree.ElementTree import ParseError
from zipfile import BadZipFile
import pytest
from openpyxl import Workbook
from openpyxl.utils.exceptions import InvalidFileException
from axelpro_riport import csv_reader as csv
from axelpro_riport import excel_files as excel


def test_missing_csv_has_filename_and_cause(tmp_path):
    path = tmp_path / "missing.csv"
    with pytest.raises(csv.CsvFileAccessError) as error:
        csv.read_movements(path)
    assert str(path) in str(error.value)
    assert isinstance(error.value.__cause__, FileNotFoundError)


@pytest.mark.parametrize("stage", ["open", "read"])
def test_csv_permission_and_mid_read_failure(tmp_path, monkeypatch, stage):
    path = tmp_path / "input.csv"

    class FailingStream:
        def __enter__(self):
            return self

        def __exit__(self, *args):
            return False

        def __iter__(self):
            return self

        def __next__(self):
            raise OSError("read failed")

    def open_file(*args, **kwargs):
        if stage == "open":
            raise PermissionError("open denied")
        return FailingStream()

    monkeypatch.setattr(Path, "open", open_file)
    with pytest.raises(csv.CsvFileAccessError) as error:
        csv.read_movements(path)
    assert isinstance(error.value.__cause__, OSError)


def test_workbook_existence_error(tmp_path, monkeypatch):
    def deny(*args):
        raise PermissionError("stat denied")

    monkeypatch.setattr(Path, "exists", deny)
    with pytest.raises(excel.WorkbookFileAccessError, match="elérhetősége"):
        excel.load_year_workbook(tmp_path, 2026, {1}, None)


@pytest.mark.parametrize(
    "failure", [PermissionError("denied"), FileNotFoundError("vanished")]
)
def test_workbook_open_access_failure(tmp_path, monkeypatch, failure):
    (tmp_path / "Napi_értékesítési_riport_2026.xlsx").touch()

    def fail(*args, **kwargs):
        raise failure

    monkeypatch.setattr(excel, "load_workbook", fail)
    with pytest.raises(excel.WorkbookFileAccessError) as error:
        excel.load_year_workbook(tmp_path, 2026, {1}, None)
    assert error.value.__cause__ is failure


@pytest.mark.parametrize(
    "failure",
    [BadZipFile("bad zip"), ParseError("bad xml"), InvalidFileException("bad xlsx")],
)
def test_corrupt_workbook_is_not_access_error(tmp_path, monkeypatch, failure):
    (tmp_path / "Napi_értékesítési_riport_2026.xlsx").touch()

    def fail(*args, **kwargs):
        raise failure

    monkeypatch.setattr(excel, "load_workbook", fail)
    with pytest.raises(excel.WorkbookValidationError) as error:
        excel.load_year_workbook(tmp_path, 2026, {1}, None)
    assert error.value.__cause__ is failure


def test_actual_non_zip_xlsx_rejected(tmp_path):
    (tmp_path / "Napi_értékesítési_riport_2026.xlsx").write_text("not an xlsx")
    with pytest.raises(excel.WorkbookValidationError):
        excel.load_year_workbook(tmp_path, 2026, {1}, None)


@pytest.mark.parametrize("stage", ["mkdir", "temporary", "write", "replace"])
def test_workbook_save_failure_preserves_target_and_cleans_temp(
    tmp_path, monkeypatch, stage
):
    path = tmp_path / "report.xlsx"
    path.write_bytes(b"original")
    wb = Workbook()
    failure = OSError("disk or permission failure")

    def fail(*args, **kwargs):
        raise failure

    if stage == "mkdir":
        monkeypatch.setattr(Path, "mkdir", fail)
    elif stage == "temporary":
        monkeypatch.setattr(excel, "NamedTemporaryFile", fail)
    elif stage == "write":
        monkeypatch.setattr(wb, "save", fail)
    else:
        monkeypatch.setattr(Path, "replace", fail)
    with pytest.raises(excel.WorkbookSaveError) as error:
        excel.save_workbook(path, wb)
    assert error.value.__cause__ is failure
    assert path.read_bytes() == b"original"
    assert list(tmp_path.iterdir()) == [path]


def test_workbook_cleanup_failure_preserves_original_error(tmp_path, monkeypatch):
    failure = OSError("write failure")
    wb = Workbook()

    def write(*args, **kwargs):
        raise failure

    def unlink(*args, **kwargs):
        raise PermissionError("cleanup failure")

    monkeypatch.setattr(wb, "save", write)
    monkeypatch.setattr(Path, "unlink", unlink)
    with pytest.raises(excel.WorkbookSaveError) as error:
        excel.save_workbook(tmp_path / "report.xlsx", wb)
    assert error.value.__cause__ is failure
