from openpyxl import Workbook, load_workbook
from sales_report.excel_files import load_year_workbook, save_workbook


def test_missing_year_creates_full_report_without_saving(tmp_path, template_sheet):
    wb, path = load_year_workbook(tmp_path, 2026, {1}, template_sheet)
    try:
        from sales_report.excel_constants import MONTH_NAMES
        assert all(name in wb.sheetnames for name in MONTH_NAMES)
        assert path.name == "Sales_Report_2026.xlsx"
        assert not path.exists()
    finally:
        wb.close()


def test_save_replaces_existing_file_without_backup_or_temp_remains(tmp_path):
    path = tmp_path / "nested" / "report.xlsx"
    wb = Workbook()
    wb.active["A1"] = "old"
    save_workbook(path, wb)
    wb.active["A1"] = "new"
    assert save_workbook(path, wb) is None
    assert load_workbook(path).active["A1"].value == "new"
    assert list(path.parent.iterdir()) == [path]


def test_existing_year_loaded_with_selected_months(tmp_path, year_report_bytes, template_sheet):
    path = tmp_path / "Sales_Report_2026.xlsx"
    path.write_bytes(year_report_bytes)
    loaded, result_path = load_year_workbook(tmp_path, 2026, {1}, template_sheet)
    assert result_path == path
    assert loaded["January"]["B2"].value.date().isoformat() == "2026-01-01"
    loaded.close()
