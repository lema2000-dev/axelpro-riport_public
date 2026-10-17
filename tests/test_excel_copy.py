import pytest
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from sales_report.excel_copy import copy_block


@pytest.mark.parametrize(
    "values,formulas", [(False, False), (True, False), (False, True), (True, True)]
)
def test_copy_translates_formulas_and_independently_clears(values, formulas):
    wb = Workbook()
    source = wb.active
    target = wb.create_sheet()
    source["B2"] = 42
    source["C2"] = "=B2+$A$1"
    source["B3"] = "Label"
    source["B2"].font = Font(bold=True)
    source["B2"].fill = PatternFill("solid", fgColor="FF0000")
    source["B2"].number_format = "0.00"
    source.row_dimensions[2].height = 30
    source.column_dimensions["B"].width = 22
    source.merge_cells("B3:C3")
    copy_block(
        source, target, 2, 3, 2, 3, 5, 5, clear_values=values, clear_formulas=formulas
    )
    assert target["E5"].value == (None if values else 42)
    assert target["F5"].value == (None if formulas else "=E5+$A$1")
    assert target["E5"].font.bold and target["E5"].fill.fgColor.rgb == "00FF0000"
    assert target["E5"].number_format == "0.00"
    assert target.row_dimensions[5].height == 30
    assert target.column_dimensions["E"].width == 22
    assert "E6:F6" in target.merged_cells
    assert source["B2"].value == 42 and source["C2"].value == "=B2+$A$1"


def test_clearing_formula_before_upward_translation_avoids_invalid_reference():
    wb = Workbook()
    target = wb.create_sheet()
    wb.active["B43"] = "=B1"
    copy_block(wb.active, target, 43, 43, 2, 2, 1, 2, clear_formulas=True)
    assert target["B1"].value is None
