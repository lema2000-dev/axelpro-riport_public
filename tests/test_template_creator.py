from copy import copy
from datetime import datetime
from openpyxl import load_workbook
from axelpro_riport.template_creator import create_template_workbook


def test_template_has_reference_structure_values_and_formulas():
    workbook = create_template_workbook()
    assert workbook.sheetnames == ["Szeptember"]
    sheet = workbook.active
    assert len(sheet.merged_cells.ranges) == 248
    assert sheet["A4"].value == "Bontott alkatrész sz.n.:"
    assert sheet["B2"].value == datetime(2026, 8, 31)
    assert sheet["N4"].value == "=B4+D4+F4+L4+H4+J4"
    assert sheet["B81"].value == "=SUM(B78:B80)"
    assert sheet["B85"].value == "=(B75-6000000)*0.05"
    assert sheet["B82"].value == 0
    assert sheet["B8"].value is None
    assert sheet.column_dimensions["A"].width == 44.25
    assert sheet.row_dimensions[34].height == 16.4


def test_template_workbooks_are_independent_and_styles_survive_save(tmp_path):
    first = create_template_workbook()
    second = create_template_workbook()
    first.active["A4"] = "Modified"
    assert second.active["A4"].value == "Bontott alkatrész sz.n.:"
    path = tmp_path / "template.xlsx"
    second.save(path)
    loaded = load_workbook(path)
    assert loaded.active["N4"].value == second.active["N4"].value
    assert copy(loaded.active["B8"].fill) == copy(second.active["B8"].fill)
    assert copy(loaded.active["B8"].border) == copy(second.active["B8"].border)
    assert set(map(str, loaded.active.merged_cells.ranges)) == set(
        map(str, second.active.merged_cells.ranges)
    )
