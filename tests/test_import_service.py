from datetime import date
from decimal import Decimal
import pytest
from openpyxl import load_workbook
from sales_report import import_service as service
from sales_report import settings
from sales_report.models import DailyExchangeRate
from sales_report.excel_constants import (
    PROCESSED_DOCUMENTS_SHEET,
)


@pytest.mark.parametrize(
    "decisions,expected",
    [(["y", "y"], [2025, 2026]), (["n", "y"], [2026]), (["n", "n"], [])],
)
def test_real_csv_to_multiple_years_only_accepted_years_saved(
    tmp_path,
    settings_location,
    write_csv,
    movement,
    monkeypatch,
    capsys,
    decisions,
    expected,
):
    directory = tmp_path / "reports"
    directory.mkdir()
    settings.create_settings(str(directory), ["Filters"])
    rows = [
        movement(day=date(2026, 1, 5), document="2026", amount="20", currency="EUR"),
        movement(day=date(2025, 1, 6), document="2025", amount="10"),
    ]
    csv_path = write_csv(tmp_path / "input.csv", rows)
    today = date.today()
    monkeypatch.setattr(
        service,
        "get_exchange_rates",
        lambda **kwargs: {
            today: DailyExchangeRate(today, today, Decimal(400)),
            date(2026, 1, 5): DailyExchangeRate(
                date(2026, 1, 5), date(2026, 1, 5), Decimal(400)
            ),
        },
    )
    responses = iter(decisions)
    monkeypatch.setattr("builtins.input", lambda prompt: next(responses))
    service.import_csv(csv_path)
    preview = capsys.readouterr().out
    assert "2025 import preview" in preview
    assert "2026 import preview" in preview
    assert "new_parts: 8,000" in preview
    assert sorted(p.name for p in directory.iterdir()) == [
        f"Sales_Report_{year}.xlsx" for year in expected
    ]
    for year in expected:
        wb = load_workbook(directory / f"Sales_Report_{year}.xlsx")
        sheet = wb["January"]
        day = rows[0].movement_date if year == 2026 else rows[1].movement_date
        column = next(
            c.column
            for c in sheet[2]
            if getattr(c.value, "date", lambda: None)() == day
        )
        assert sheet.cell(9, column).value == (8000 if year == 2026 else 10)
        assert list(wb[PROCESSED_DOCUMENTS_SHEET].values)[1:] == [(year, str(year))]
    if expected:
        # Re-import the same CSV: accepted files retain their amounts and keys.
        responses = iter("y" if year in expected else "n" for year in (2025, 2026))
        monkeypatch.setattr("builtins.input", lambda prompt: next(responses))
        service.import_csv(csv_path)
        preview = capsys.readouterr().out
        if expected == [2025, 2026]:
            assert "There are no new items to import." in preview
            assert "import preview" not in preview
        else:
            assert "Previously processed items skipped: 1" in preview
        for year in expected:
            wb = load_workbook(directory / f"Sales_Report_{year}.xlsx")
            sheet = wb["January"]
            day = rows[0].movement_date if year == 2026 else rows[1].movement_date
            column = next(
                c.column
                for c in sheet[2]
                if getattr(c.value, "date", lambda: None)() == day
            )
            assert sheet.cell(9, column).value == (8000 if year == 2026 else 10)
            assert list(wb[PROCESSED_DOCUMENTS_SHEET].values)[1:] == [(year, str(year))]
