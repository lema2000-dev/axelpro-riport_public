"""Test workbook validation."""

from datetime import date
from unittest.mock import Mock
import warnings
import pytest
from openpyxl import load_workbook
from openpyxl.cell.cell import MergedCell
from openpyxl.styles import Font
from sales_report import excel_files as f
from sales_report.excel_constants import MONTH_NAMES


@pytest.fixture
def report_path(tmp_path, year_report_bytes):
    path = tmp_path / "Sales_Report_2026.xlsx"
    path.write_bytes(year_report_bytes)
    return path



def test_valid_report_loads_with_imported_and_manual_values(
    report_path
):
    wb = load_workbook(report_path)
    wb["January"]["J8"] = 120
    wb["January"]["B74"] = 20.5
    wb.save(report_path)
    loaded, path = f.load_year_workbook(
        report_path.parent, 2026, set(range(1, 13)), None
    )
    assert path == report_path
    assert loaded["January"]["J8"].value == 120
    assert loaded["January"]["B74"].value == 20.5
    loaded.close()


def test_missing_sheet_is_reported_before_headers(report_path):
    wb = load_workbook(report_path)
    del wb["February"]
    wb["January"]["B2"] = "also bad"
    wb.save(report_path)
    with pytest.raises(f.WorkbookValidationError) as error:
        f.load_year_workbook(report_path.parent, 2026, set(range(1, 13)), None)
    assert "February" in str(error.value) and "Missing" in str(error.value)


@pytest.mark.parametrize(
    "cell,value",
    [
        ("B1", "Wrong"),
        ("B2", date(2026, 1, 2)),
        ("H2", 99),
        ("B44", "Wrong month"),
        ("A2", "Unexpected"),
    ],
)
def test_month_headers_report_exact_location(report_path, cell, value):
    wb = load_workbook(report_path)
    sheet = wb["January"]
    if isinstance(sheet[cell], MergedCell):
        for merged_range in list(sheet.merged_cells.ranges):
            if cell in merged_range:
                sheet.unmerge_cells(str(merged_range))
                break
    sheet[cell] = value
    wb.save(report_path)
    with pytest.raises(f.WorkbookValidationError) as error:
        f.load_year_workbook(report_path.parent, 2026, set(range(1, 13)), None)
    assert "January" in str(error.value) and cell in str(error.value)
    assert str(report_path) in str(error.value)


@pytest.mark.parametrize(
    "sheet,cell",
    [
        ("Summary", "A1"),
        ("Summary", "M2"),
        ("Exchange_Rates", "A13"),
        ("_daily_exchange_rates", "B1"),
        ("_imported_documents", "A1"),
    ],
)
def test_other_sheet_headers_report_exact_location(report_path, sheet, cell):
    wb = load_workbook(report_path)
    wb[sheet][cell] = "Wrong"
    wb.save(report_path)
    with pytest.raises(f.WorkbookValidationError) as error:
        f.load_year_workbook(report_path.parent, 2026, set(range(1, 13)), None)
    assert sheet in str(error.value) and cell in str(error.value)


@pytest.mark.parametrize(
    "action,expected", [("remove", "Missing"), ("add", "Unexpected")]
)
def test_merge_errors(year_report, reference_report, tmp_path, action, expected):
    sheet = year_report["January"]
    if action == "remove":
        sheet.unmerge_cells("B8:C8")
    else:
        sheet.merge_cells("P90:Q90")
    with pytest.raises(f.WorkbookValidationError, match=expected):
        f._validate_merged_ranges(
            year_report, reference_report, tmp_path / "report.xlsx"
        )


@pytest.mark.parametrize(
    "cell,value,description",
    [("A8", "wrong", "label"), ("J12", "=1+2", "formula"), ("J12", 123, "formula")],
)
def test_fixed_label_and_formula_validation(
    year_report, reference_report, tmp_path, cell, value, description
):
    year_report["January"][cell] = value
    with pytest.raises(f.WorkbookValidationError) as error:
        f._validate_labels_and_formulas(
            year_report, reference_report, tmp_path / "report.xlsx", list(MONTH_NAMES)
        )
    assert cell in str(error.value) and description in str(error.value)


def test_unchanged_labels_and_formulas_are_valid(
    year_report, reference_report, tmp_path
):
    f._validate_labels_and_formulas(
        year_report, reference_report, tmp_path / "report.xlsx", list(MONTH_NAMES)
    )


@pytest.mark.parametrize("cell", ["B7", "A49", "B42"])
def test_reserved_empty_cells(year_report, reference_report, tmp_path, cell):
    year_report["January"][cell] = "unexpected"
    with pytest.raises(f.WorkbookValidationError) as error:
        f._validate_empty_cells(year_report, reference_report, tmp_path / "report.xlsx", list(MONTH_NAMES))
    assert cell in str(error.value)


@pytest.mark.parametrize(
    "value", ["text", True, "=1+2", date(2026, 1, 5), float("inf")]
)
def test_data_cells_reject_non_numeric_values(
    year_report, reference_report, tmp_path, value
):
    year_report["January"]["J8"] = value
    with pytest.raises(f.WorkbookValidationError, match="J8"):
        f._validate_numeric_cells(
            year_report, reference_report, tmp_path / "report.xlsx", list(MONTH_NAMES)
        )


@pytest.mark.parametrize("value", [None, 0, -12, 2.5])
def test_data_cells_accept_empty_and_finite_numbers(
    year_report, reference_report, tmp_path, value
):
    year_report["January"]["J8"] = value
    f._validate_numeric_cells(year_report, reference_report, tmp_path / "report.xlsx", list(MONTH_NAMES))


@pytest.mark.parametrize(
    "kind,description",
    [("row", "row height"), ("column", "column width"), ("font", "font")],
)
def test_formatting_differences_warn_without_interrupting(
    year_report, reference_report, tmp_path, kind, description
):
    sheet = year_report["January"]
    if kind == "row":
        sheet.row_dimensions[8].height = 99
    elif kind == "column":
        sheet.column_dimensions["B"].width = 99
    else:
        sheet["J8"].font = Font(name="Different", size=30)
    with warnings.catch_warnings(record=True) as collected:
        warnings.simplefilter("always", f.WorkbookValidationWarning)
        f._validate_formatting_and_dimensions(
            year_report, reference_report, tmp_path / "report.xlsx", list(MONTH_NAMES)
        )
    assert any(description in str(w.message) for w in collected)


def test_failed_validation_closes_loaded_workbook(report_path, monkeypatch):
    wb = load_workbook(report_path)
    wb.close = Mock(wraps=wb.close)
    del wb["January"]
    monkeypatch.setattr(f, "load_workbook", lambda *args, **kwargs: wb)
    with pytest.raises(f.WorkbookValidationError):
        f.load_year_workbook(report_path.parent, 2026, set(range(1, 13)), None)
    wb.close.assert_called_once()


def test_failed_late_validation_closes_report_and_reference(
    report_path, year_report, reference_report, monkeypatch
):
    year_report.close = Mock(wraps=year_report.close)
    reference_report.close = Mock(wraps=reference_report.close)
    monkeypatch.setattr(f, "load_workbook", lambda *args, **kwargs: year_report)
    monkeypatch.setattr(f, "create_partial_year_workbook", lambda **kwargs: reference_report)

    def fail(*args, **kwargs):
        raise f.WorkbookValidationError("bad merge")

    monkeypatch.setattr(f, "_validate_merged_ranges", fail)
    with pytest.raises(f.WorkbookValidationError, match="bad merge"):
        f.load_year_workbook(report_path.parent, 2026, set(range(1, 13)), None)
    year_report.close.assert_called_once()
    reference_report.close.assert_called_once()
